from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# Create presentation
prs = Presentation()

# Constants
BG_COLOR = RGBColor(18, 18, 18)  # Deep Charcoal
ACCENT_COLOR = RGBColor(0, 212, 255)  # Electric Blue
WHITE = RGBColor(255, 255, 255)
GRAY = RGBColor(150, 150, 150)

def apply_base_styling(slide):
    """Applies background color and thin circuit lines to a slide."""
    # Background
    background = slide.shapes.add_textbox(0, 0, prs.slide_width, prs.slide_height)
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = BG_COLOR
    
    # Subtle PCB-style lines (corner accents)
    line1 = slide.shapes.add_connector(1, Inches(0.5), Inches(0.5), Inches(2), Inches(0.5))
    line1.line.color.rgb = ACCENT_COLOR
    line1.line.width = Pt(0.5)
    
    line2 = slide.shapes.add_connector(1, Inches(0.5), Inches(0.5), Inches(0.5), Inches(1.5))
    line2.line.color.rgb = ACCENT_COLOR
    line2.line.width = Pt(0.5)

def add_content_slide(title_text, body_items, visual_desc=""):
    slide_layout = prs.slide_layouts[1] # Title and Content
    slide = prs.slides.add_slide(slide_layout)
    apply_base_styling(slide)
    
    # Title
    title = slide.shapes.title
    title.text = title_text
    title_text_frame = title.text_frame.paragraphs[0]
    title_text_frame.font.color.rgb = WHITE
    title_text_frame.font.size = Pt(36)
    title_text_frame.font.bold = True
    
    # Body
    if body_items:
        body_shape = slide.placeholders[1]
        tf = body_shape.text_frame
        tf.word_wrap = True
        for item in body_items:
            p = tf.add_paragraph()
            p.text = f"• {item}"
            p.font.color.rgb = WHITE
            p.font.size = Pt(18)
            p.space_after = Pt(10)
            
    # Visual Placeholder
    if visual_desc:
        box = slide.shapes.add_textbox(Inches(5.5), Inches(1.5), Inches(4), Inches(5))
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.add_paragraph()
        p.text = f"[TECHNICAL VISUAL: {visual_desc}]"
        p.font.color.rgb = ACCENT_COLOR
        p.font.size = Pt(12)
        p.font.italic = True
        p.alignment = PP_ALIGN.CENTER

# --- SLIDE 1: TITLE ---
slide_1 = prs.slides.add_slide(prs.slide_layouts[0])
apply_base_styling(slide_1)
title = slide_1.shapes.title
title.text = "From Algorithm to Silicon"
title.text_frame.paragraphs[0].font.color.rgb = WHITE
title.text_frame.paragraphs[0].font.size = Pt(44)

subtitle = slide_1.placeholders[1]
subtitle.text = "Design & Fabrication of a Small Custom Microchip\n[ Team Name / Project Code ]"
subtitle.text_frame.paragraphs[0].font.color.rgb = ACCENT_COLOR

# --- SLIDE 2: THE PROBLEM ---
add_content_slide(
    "The Gap: Theory vs. Physics",
    [
        "High-level algorithms ignore physical constraints.",
        "CPUs/GPUs are jacks-of-all-trades, masters of none.",
        "Silicon real-estate requires precise optimization of power, timing, and area (PPA).",
        "The 'Abstraction Gap': Python code doesn't account for gate-level leakage or wire-load delays."
    ],
    "Split-screen diagram. Left: Python code snippet. Right: SEM image of a transistor or dense metal layers."
)

# --- SLIDE 3: OUR IDEA ---
add_content_slide(
    "The Vision: A Full ASIC Journey",
    [
        "Goal: Design, verify, and tape-out a custom IC.",
        "Scope: Full-stack hardware engineering from mathematical model to GDSII.",
        "Current Focus: Architecture Exploration and Feasibility Analysis.",
        "Preparing a design ready for Multi-Project Wafer (MPW) fabrication."
    ],
    "Circular flow: Golden Model → RTL → Physical Layout → Silicon Die."
)

# --- SLIDE 4: WHAT ARE WE EXPLORING? ---
add_content_slide(
    "Architecture Candidates",
    [
        "Candidate A: MAC Accelerator (Fixed-point DSP).",
        "Candidate B: Bit-Serial ALU (Ultra-low power).",
        "Candidate C: TinyML Core (Minimalist NN inference).",
        "Candidate D: Cryptographic Primitive (AES/SHA block).",
        "Note: Final selection depends on PPA benchmarking."
    ],
    "Four-quadrant grid of simplified block diagrams showing candidate datapath architectures."
)

# --- SLIDE 5: HOW WILL WE CHOOSE? ---
add_content_slide(
    "Selection & Feasibility",
    [
        "Area: Must fit within 180nm/65nm MPW die constraints.",
        "Power: Analysis of switching activity and leakage.",
        "Complexity: Verification cycles required for 100% coverage.",
        "Foundry Support: Compatibility with standard cell libraries."
    ],
    "Radar chart (Spider Web) showing trade-offs between Area, Power, Complexity, and Relevance."
)

# --- SLIDE 6: ALGORITHM TO HARDWARE ---
add_content_slide(
    "The Logical Transformation",
    [
        "Golden Model: Algorithm in Python/C++.",
        "RTL Design: Verilog/SystemVerilog implementation.",
        "Functional Verification: Testbenches and Logic Simulation.",
        "Transition: Mapping mathematical logic to flip-flops and gates."
    ],
    "Linear flowchart: Algorithm → Block Diagram → Verilog Code → Simulation Waveforms."
)

# --- SLIDE 7: ASIC JOURNEY ---
add_content_slide(
    "Physical Design Pipeline",
    [
        "Synthesis: RTL logic to gate-level netlist.",
        "Floorplanning & Placement: Strategic cell positioning.",
        "Clock Tree Synthesis (CTS): Balancing signal skew.",
        "Routing & Sign-off: Metal layer connectivity and DRC/LVS."
    ],
    "Chevron-style pipeline showing the RTL-to-GDSII industrial flow."
)

# --- SLIDE 8: FINAL DELIVERABLE ---
add_content_slide(
    "Tape-Out Package",
    [
        "GDSII Database: The physical 'mask' for the foundry.",
        "Synthesizable RTL: Verified and documented.",
        "Timing/Power Reports: Guaranteed performance sign-off.",
        "Post-Silicon Test Plan: BIST and scan-chain strategy."
    ],
    "High-contrast zoom of a GDSII layout heatmap showing dense wiring and standard cells."
)

# --- SLIDE 9: PROJECT ROADMAP ---
add_content_slide(
    "Milestones & Current Status",
    [
        "Q1: Exploration & Selection [YOU ARE HERE].",
        "Q2: RTL Implementation & Functional Verification.",
        "Q3: Physical Design (Synthesis, P&R).",
        "Q4: Sign-off & Tape-out."
    ],
    "Horizontal timeline with an Electric Blue marker on the 'Exploration' phase."
)

# --- SLIDE 10: CLOSING VISION ---
slide_10 = prs.slides.add_slide(prs.slide_layouts[1])
apply_base_styling(slide_10)
title = slide_10.shapes.title
title.text = "From Algorithm to Silicon"
title.text_frame.paragraphs[0].font.color.rgb = ACCENT_COLOR

body = slide_10.placeholders[1]
body.text = "“Our goal is to take a small computational idea beyond simulation — and turn it into silicon.”\n\nThank You. Questions?"
p = body.text_frame.paragraphs[0]
p.font.size = Pt(32)
p.font.bold = True
p.alignment = PP_ALIGN.CENTER
p.font.color.rgb = WHITE

# Save
prs.save('/mnt/data/From_Algorithm_to_Silicon.pptx')